from __future__ import annotations

import csv
import shutil
import sqlite3
from pathlib import Path

from openpyxl import Workbook

from src.models import FIELD_LABELS, JobPosting


def write_outputs(
    jobs: list[JobPosting],
    output_dir: str | Path,
    run_date: str,
    write_latest: bool = False,
) -> dict[str, Path]:
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    sorted_jobs = sorted(jobs, key=lambda job: job.total_score, reverse=True)
    base_name = f"{run_date}_实习信息"
    paths = {
        "xlsx": target_dir / f"{base_name}.xlsx",
        "csv": target_dir / f"{base_name}.csv",
        "db": target_dir / f"{base_name}.db",
    }
    write_xlsx(sorted_jobs, paths["xlsx"])
    write_csv(sorted_jobs, paths["csv"])
    write_sqlite(sorted_jobs, paths["db"])
    if write_latest:
        shutil.copyfile(paths["xlsx"], target_dir / "jobs_latest.xlsx")
        shutil.copyfile(paths["csv"], target_dir / "jobs_latest.csv")
        shutil.copyfile(paths["db"], target_dir / "jobs.db")
    return paths


def write_xlsx(jobs: list[JobPosting], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "实习信息"
    headers = [label for _, label in FIELD_LABELS]
    sheet.append(headers)
    for job in jobs:
        row = job.to_output_row()
        sheet.append([row[header] for header in headers])
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 40)
    workbook.save(path)


def write_csv(jobs: list[JobPosting], path: Path) -> None:
    headers = [label for _, label in FIELD_LABELS]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for job in jobs:
            writer.writerow(job.to_output_row())


def write_sqlite(jobs: list[JobPosting], path: Path) -> None:
    if path.exists():
        path.unlink()
    headers = [label for _, label in FIELD_LABELS]
    with sqlite3.connect(path) as connection:
        columns = ", ".join(f'"{header}" TEXT' for header in headers)
        connection.execute(f'CREATE TABLE jobs (id INTEGER PRIMARY KEY AUTOINCREMENT, {columns})')
        placeholders = ", ".join("?" for _ in headers)
        quoted_headers = ", ".join(f'"{header}"' for header in headers)
        for job in jobs:
            row = job.to_output_row()
            values = [str(row[header]) for header in headers]
            connection.execute(f"INSERT INTO jobs ({quoted_headers}) VALUES ({placeholders})", values)
        connection.commit()

